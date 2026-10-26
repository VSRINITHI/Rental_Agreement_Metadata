"""Document loading: DOCX via python-docx, scanned PNG via NVIDIA Nemotron Parse.

No extraction logic lives here -- only conversion of a file into plain text.
"""
from __future__ import annotations

import json
from pathlib import Path

import docx
import requests
from docx.table import Table
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

import config

_COORD_PREFIXES = ("x_", "y_")
_CLASS_PREFIX = "class_"


def _is_layout_tag(inner: str) -> bool:
    """True for Nemotron coordinate/class markers such as 'x_0.12', 'y_0.5', 'class_Text'."""
    if inner.startswith(_CLASS_PREFIX):
        return len(inner) > len(_CLASS_PREFIX)
    if inner[:2] in _COORD_PREFIXES:
        number = inner[2:]
        return number.replace(".", "", 1).isdigit()
    return False


def strip_layout_tags(text: str) -> str:
    """Remove OCR layout markers (bounding-box coordinates / block classes).

    This is OCR-markup cleanup, not field extraction. Any other markup is kept.
    """
    out: list[str] = []
    i, n = 0, len(text)
    while i < n:
        ch = text[i]
        if ch == "<":
            close = text.find(">", i + 1)
            if close != -1 and _is_layout_tag(text[i + 1 : close]):
                left = out[-1][-1:] if out else ""
                j = close + 1
                # swallow directly adjacent layout tags as one cluster
                while j < n and text[j] == "<":
                    c2 = text.find(">", j + 1)
                    if c2 == -1 or not _is_layout_tag(text[j + 1 : c2]):
                        break
                    j = c2 + 1
                right = text[j : j + 1]
                if left.isalnum() and right.isalnum():
                    out.append(" ")
                i = j
                continue
        out.append(ch)
        i += 1
    lines = [" ".join(line.split()) for line in "".join(out).split("\n")]
    cleaned: list[str] = []
    for line in lines:
        if line or (cleaned and cleaned[-1]):
            cleaned.append(line)
    return "\n".join(cleaned).strip()


def _iter_body(document: docx.document.Document):
    """Yield paragraphs and tables in true document order."""
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _table_lines(table: Table) -> list[str]:
    lines: list[str] = []
    for row in table.rows:
        seen_cells: set[int] = set()
        cells: list[str] = []
        for cell in row.cells:
            if id(cell._tc) in seen_cells:  # merged cells repeat the same element
                continue
            seen_cells.add(id(cell._tc))
            text = " ".join(cell.text.split())
            if text:
                cells.append(text)
        if cells:
            lines.append(" | ".join(cells))
    return lines


def load_docx(path: Path) -> str:
    """Text of a DOCX: header, body (paragraphs and tables in order), footer."""
    try:
        document = docx.Document(str(path))
    except Exception as exc:  # corrupt/unsupported file
        raise RuntimeError(f"Cannot open DOCX '{path.name}': {exc}") from exc

    blocks: list[str] = []
    for section in document.sections:
        blocks += [p.text.strip() for p in section.header.paragraphs if p.text.strip()]
    for item in _iter_body(document):
        if isinstance(item, Paragraph):
            if item.text.strip():
                blocks.append(item.text.strip())
        else:
            blocks += _table_lines(item)
    for section in document.sections:
        blocks += [p.text.strip() for p in section.footer.paragraphs if p.text.strip()]

    text = "\n".join(blocks).strip()
    if not text:
        raise RuntimeError(f"DOCX '{path.name}' contains no readable text.")
    return text


def _error_detail(response: requests.Response) -> str:
    try:
        body = response.json()
        return str(body.get("error") or body.get("message") or body.get("detail") or "")
    except Exception:
        return response.text[:200].strip()


def _parse_response_text(body: object) -> str:
    if isinstance(body, dict):
        if "content" in body:
            return str(body["content"])
        choices = body.get("choices")
        if isinstance(choices, list) and choices:
            first = choices[0]
            return str(first.get("message", {}).get("content", first))
    return json.dumps(body)


def load_png(path: Path, settings: config.Settings | None = None, session=None) -> str:
    """OCR/parse a scanned image with NVIDIA Nemotron Parse."""
    settings = settings or config.get_settings()
    api_key = config.require_api_key(settings)
    session = session or requests
    try:
        with open(path, "rb") as fh:
            image = fh.read()
    except OSError as exc:
        raise RuntimeError(f"Cannot read image '{path.name}': {exc}") from exc

    last_error = "unknown error"
    for attempt in range(1, settings.max_retries + 1):
        try:
            response = session.post(
                settings.parse_url,
                headers={"Authorization": f"Bearer {api_key}", "Accept": "application/json"},
                files={"document": (path.name, image, "image/png")},
                data={"options": json.dumps({"decoding": {"max_output_tokens": 4096}})},
                timeout=(settings.connect_timeout, settings.read_timeout),
            )
        except (requests.Timeout, requests.ConnectionError) as exc:
            last_error = type(exc).__name__
        else:
            if response.status_code == 200:
                return strip_layout_tags(_parse_response_text(response.json()))
            last_error = f"HTTP {response.status_code}: {_error_detail(response)}"
            if response.status_code not in (408, 429, 500, 502, 503, 504):
                break
        if attempt < settings.max_retries:
            import time

            time.sleep(min(2**attempt, 20))
    raise RuntimeError(f"Nemotron Parse failed for '{path.name}': {last_error}")


def load_document(path: Path | str, settings: config.Settings | None = None) -> str:
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"Document not found: {path}")
    name = path.name.lower()
    if name.endswith(".docx"):
        return load_docx(path)
    if name.endswith(".png"):
        return load_png(path, settings)
    raise ValueError(f"Unsupported format '{path.name}'. Only .docx and .png are supported.")
